                       
from __future__ import annotations


                                    
DATA_YAML = r""
MODEL_TAG = "YOLO11m_OPT_105_960_Trial8_train_then_external"
OPT_MODEL = True
EPOCHS = 300
IMGSZ = 960
BATCH = 10
DEVICE = "0"
WORKERS = 0
SEED = 42
AMP = True
SAVE_PERIOD = 25
OPTIMIZER = "SGD"
USE_PRETRAINED = False  
USE_WIOU = True 
EVALUATE_TEST_AFTER_TRAINING = False                       
                                                       
                                                     
RUN_MODE = "train_then_external"                                
                                                                 
EXISTING_BEST_PT = r""                    
                                           
EXTERNAL_IMAGES = r""
EXTERNAL_LABELS = r""
                        
SAVE_EXTERNAL_VISUALS = True
                                   
ALLOW_MISSING_EXTERNAL_LABELS_AS_EMPTY = False
EXPORT_CONF = 0.01
CM_CONF = 0.25
NMS_IOU = 0.70
MATCH_IOU = 0.50
TRIAL8_BEST_HYP = {
    "lr0": 0.0184447, "lrf": 0.1620807,
    "momentum": 0.8710642, "weight_decay": 0.0001,
    "warmup_epochs": 3.6216454, "hsv_h": 0.0,
    "hsv_s": 0.5339505, "hsv_v": 0.2780041,
    "translate": 0.1706818, "scale": 0.4860711,
    "fliplr": 0.3069651,
}

import csv
import datetime as dt
import hashlib
import json
import math
import os
import shutil
import sys
import traceback
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parent
IMAGE_EXT = {'.jpg', '.jpeg', '.bmp', '.png', '.tiff', '.tif', '.webp'}


def json_save(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda: f.read(1024 * 1024), b''):
            h.update(b)
    return h.hexdigest()


def csv_save(path, rows, columns=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if columns is None:
        columns = list(rows[0]) if rows else ['note']
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(columns), extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)


def xlsx_save(path, sheets):
                                           
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        cols = list(cols)
        ws.append(cols)
        for r in rows:
            vals = []
            for c in cols:
                v = r.get(c)
                if isinstance(v, np.generic):
                    v = v.item()
                if isinstance(v, (list, dict, tuple)):
                    v = json.dumps(v, ensure_ascii=False)
                if isinstance(v, float) and not math.isfinite(v):
                    v = None
                vals.append(v)
            ws.append(vals)
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='174A76')
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center')
        for colidx, col in enumerate(cols, 1):
            width = 18
            if 'path' in col.lower() or 'image' in col.lower() or 'file' in col.lower():
                width = 38
            elif col.lower().startswith('pred_'):
                width = 15
            ws.column_dimensions[get_column_letter(colidx)].width = width
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def tabular_save(stem, rows, columns=None):
    rows = list(rows)
    cols = list(columns) if columns else (list(rows[0]) if rows else ['note'])
    csv_save(str(stem) + '.csv', rows, cols)
    xlsx_save(str(stem) + '.xlsx', {'Data': (cols, rows)})


def configure_plot():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        'font.family': 'Arial', 'font.size': 10, 'axes.labelsize': 12,
        'axes.spines.top': False, 'axes.spines.right': False,
        'axes.grid': False, 'figure.facecolor': 'white',
        'savefig.facecolor': 'white', 'pdf.fonttype': 42, 'ps.fonttype': 42,
    })
    return plt


def save_figure(fig, stem):
    stem = Path(stem)
    stem.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(str(stem) + '.png', dpi=300, facecolor='white')
    fig.savefig(str(stem) + '.tiff', dpi=600, facecolor='white',
                pil_kwargs={'compression': 'tiff_lzw'})


def read_data(data_yaml):
    import yaml
    yaml_path = Path(data_yaml).expanduser().resolve()
    if not yaml_path.is_file():
        raise FileNotFoundError(f'DATA_YAML not found: {yaml_path}')
    obj = yaml.safe_load(yaml_path.read_text(encoding='utf-8-sig'))
    names = obj.get('names')
    if isinstance(names, dict):
        names = [names.get(i, names.get(str(i))) for i in range(len(names))]
    if not isinstance(names, list) or len(names) != 105:
        raise ValueError('This experiment requires a 105-class YAML with exactly 105 names')
    names = [str(s).strip() for s in names]
    if names[:38] != [str(i) for i in range(1, 39)]:
        raise ValueError('Invalid label mapping: the first 38 labels must be named 1..38, with internal IDs 0..37')
    if len(set(names)) != 105:
        raise ValueError('Duplicate class names in YAML')
    if obj.get('nc') is not None and int(obj['nc']) != 105:
        raise ValueError('YAML nc must be 105')
    root = Path(str(obj.get('path') or yaml_path.parent))
    if not root.is_absolute():
        root = yaml_path.parent / root
    root = root.resolve()
    splits = {}
    for part in ('train', 'val', 'test'):
        if part not in obj:
            if part == 'test':
                continue
            raise ValueError(f'data.yaml missing {part}')
        p = Path(obj[part])
        p = (root / p).resolve() if not p.is_absolute() else p.resolve()
        if not p.is_dir():
            raise FileNotFoundError(f'{part}image directory not found: {p}')
        splits[part] = p
    return yaml_path, names, splits


def image_to_label(path):
    path = Path(path)
    pp = list(path.parts)
    ix = [i for i, s in enumerate(pp) if s.lower() == 'images']
    if not ix:
        raise ValueError(f'Image path has no images component: {path}')
    pp[ix[-1]] = 'labels'
    return Path(*pp).with_suffix('.txt')


def gt_boxes(image):
    from PIL import Image
    with Image.open(image) as f:
        width, height = f.size
    p = image_to_label(image)
    if not p.is_file():
        raise FileNotFoundError(f'Image is missing TXT annotation: {p}')
    out = []
    for j, line in enumerate(p.read_text(encoding='utf-8-sig').splitlines(), 1):
        if not line.strip():
            continue
        parts = line.split()
        if len(parts) != 5:
            raise ValueError(f'{p}:{j}must have five columns')
        nums = [float(v) for v in parts]
        c, x, y, w, h = nums
        if not (c.is_integer() and 0 <= c < 105 and all(math.isfinite(a) for a in nums)
                and 0 <= x <= 1 and 0 <= y <= 1 and 0 < w <= 1 and 0 < h <= 1
                and x - w/2 >= -1e-4 and x + w/2 <= 1+1e-4
                and y - h/2 >= -1e-4 and y + h/2 <= 1+1e-4):
            raise ValueError(f'Invalid label coordinates or class ID: {p}:{j}')
        rect = [(x-w/2)*width, (y-h/2)*height, (x+w/2)*width, (y+h/2)*height]
        out.append({'id': int(c), 'box': rect, 'index': j})
    return out, width, height


def audit_dataset(names, split_dirs, folder):
    per_class = []
    img_rows = []
    split_total = Counter()
    tally = defaultdict(Counter)
    for split, imgdir in split_dirs.items():
        imgs = sorted(p for p in imgdir.rglob('*') if p.is_file() and p.suffix.lower() in IMAGE_EXT)
        if not imgs:
            raise ValueError(f'Dataset {split} contains no images: {imgdir}')
        for image in imgs:
            gt, width, height = gt_boxes(image)
            numeric = [g for g in gt if g['id'] < 38]
            ids = sorted(set(g['id'] for g in numeric))
            split_total[split] += 1
            for g in gt:
                tally[split][g['id']] += 1
            img_rows.append({'split': split, 'image': str(image), 'label_file': str(image_to_label(image)),
                             'width': width, 'height': height, 'objects_all': len(gt),
                             'objects_numeric': len(numeric), 'numeric_ids_0based': ','.join(map(str, ids)),
                             'paper_numbers_1based': ','.join(str(i+1) for i in ids),
                             'numeric_one_per_image': len(numeric) == 1})
    for i, name in enumerate(names):
        per_class.append({'yolo_id': i, 'display_name': name, 'paper_number': i+1 if i < 38 else None,
                          **{f'{split}_objects': tally[split][i] for split in split_dirs}})
    folder = Path(folder)
    tabular_save(folder / 'dataset_per_class', per_class)
    tabular_save(folder / 'dataset_per_image', img_rows)
    tabular_save(folder / 'class_mapping_105', [
        {'yolo_id_0based': i, 'yaml_name': name, 'is_numeric': i < 38,
         'paper_id_1based': i+1 if i < 38 else None} for i, name in enumerate(names)])
    json_save(folder / 'dataset_counts.json', dict(split_total))
    print('[DATA] Image counts by split', dict(split_total))
    return dict(split_total)


def iou(a, b):
    x0, y0 = max(a[0],b[0]), max(a[1],b[1])
    x1, y1 = min(a[2],b[2]), min(a[3],b[3])
    inter = max(0,x1-x0)*max(0,y1-y0)
    ar_a = max(0,a[2]-a[0])*max(0,a[3]-a[1])
    ar_b = max(0,b[2]-b[0])*max(0,b[3]-b[1])
    den = ar_a+ar_b-inter
    return float(inter/den) if den else 0.0


def matching(gt, pred, threshold):
                                               
    pairs = []
    for a, g in enumerate(gt):
        for b, p in enumerate(pred):
            score = iou(g['box'], p['box'])
            if score >= threshold:
                pairs.append((score, p['confidence'], a, b))
    pairs.sort(key=lambda t: (-t[0], -t[1], t[2], t[3]))
    used_gt, used_pred, matches = set(), set(), []
    for ov, cf, a, b in pairs:
        if a not in used_gt and b not in used_pred:
            matches.append((a, b, ov))
            used_gt.add(a)
            used_pred.add(b)
    return matches, used_gt, used_pred


def matrix_report(gt_records, pred_records, names, split, scope, out_dir, conf_threshold, iou_threshold):
    nc = len(names)
    matrix = np.zeros((nc,nc), dtype=np.int64)
    misses, fp, counts_gt, counts_pred = np.zeros((4,nc),dtype=np.int64)
    detail = []
    image_stats = []
    for path, gt_all in gt_records.items():
        pred_all = pred_records[path]
        gt = [g for g in gt_all if g['id'] < nc]
        pred = [p for p in pred_all if p['id'] < nc and p['confidence'] >= conf_threshold]
        for g in gt:
            counts_gt[g['id']] += 1
        for p in pred:
            counts_pred[p['id']] += 1
        matches, used_gt, used_pred = matching(gt, pred, iou_threshold)
        for a,b,overlap in matches:
            g,p = gt[a],pred[b]
            matrix[g['id'],p['id']] += 1
            detail.append({'image':path, 'status':'correct' if g['id']==p['id'] else 'misclassified',
                           'true_yolo_id':g['id'],'true_name':names[g['id']],
                           'pred_yolo_id':p['id'],'pred_name':names[p['id']],
                           'confidence':p['confidence'],'match_iou':overlap,
                           **{f'gt_{k}':v for k,v in zip(('x1','y1','x2','y2'),g['box'])},
                           **{f'pred_{k}':v for k,v in zip(('x1','y1','x2','y2'),p['box'])}})
        for a,g in enumerate(gt):
            if a not in used_gt:
                misses[g['id']] += 1
                detail.append({'image':path,'status':'missed_gt','true_yolo_id':g['id'],
                               'true_name':names[g['id']],
                               **{f'gt_{k}':v for k,v in zip(('x1','y1','x2','y2'),g['box'])}})
        for b,p in enumerate(pred):
            if b not in used_pred:
                fp[p['id']] += 1
                detail.append({'image':path,'status':'unmatched_prediction','pred_yolo_id':p['id'],
                               'pred_name':names[p['id']], 'confidence':p['confidence'],
                               **{f'pred_{k}':v for k,v in zip(('x1','y1','x2','y2'),p['box'])}})
        image_stats.append({'image':path,'gt_count':len(gt),'pred_count':len(pred),
                            'matched':len(matches),'missed':len(gt)-len(used_gt),
                            'unmatched_predictions':len(pred)-len(used_pred),
                            'correct_matches':sum(gt[a]['id']==pred[b]['id'] for a,b,_ in matches)})
    matched = matrix.sum(axis=1)
    norm = np.divide(matrix,matched[:,None],out=np.zeros_like(matrix,dtype=float),where=matched[:,None]>0)
    cols=['true_label'] + [f'pred_{label}' for label in names]
    rows_counts = [{'true_label':str(name), **{f'pred_{s}':int(matrix[i,j]) for j,s in enumerate(names)}}
                   for i,name in enumerate(names)]
    rows_norm = [{'true_label':str(name), **{f'pred_{s}':float(norm[i,j]) for j,s in enumerate(names)}}
                 for i,name in enumerate(names)]
    per_class = []
    for i,name in enumerate(names):
        c, m, n = int(matrix[i,i]), int(matched[i]), int(counts_gt[i])
        per_class.append({'yolo_id_0based':i,'paper_number_1based':i+1 if i<38 else None,
                          'class_name':name,'GT_total':n,'matched_GT':m,'correct_matches':c,
                          'wrong_class_matches':m-c,'missed_GT':int(misses[i]),
                          'predictions':int(counts_pred[i]),'unmatched_predictions':int(fp[i]),
                          'matched_classification_accuracy':c/m if m else None,
                          'correct_detection_fraction_of_GT':c/n if n else None})
    total_gt,total_p,total_m,total_corr = int(counts_gt.sum()),int(counts_pred.sum()),int(matched.sum()),int(matrix.trace())
    summary = {'split':split,'scope':scope,'classes':nc,'GT':total_gt,'predictions':total_p,
               'matched':total_m,'correct_matches':total_corr,
               'missed_GT':int(misses.sum()),'unmatched_predictions':int(fp.sum()),
               'matched_classification_accuracy':total_corr/total_m if total_m else None,
               'correct_detection_fraction_of_GT':total_corr/total_gt if total_gt else None,
               'conf_threshold':conf_threshold,'match_iou':iou_threshold,
               'note':'Classification accuracy among matched objects; misses and false positives are reported separately. This is not overall detection accuracy or mAP.'}
    prefix=Path(out_dir)/f'{split}_{scope}'
    csv_save(f'{prefix}_counts.csv',rows_counts,cols)
    csv_save(f'{prefix}_normalized.csv',rows_norm,cols)
    tabular_save(f'{prefix}_per_class',per_class)
    detail_cols=list(dict.fromkeys(k for row in detail for k in row)) if detail else ['image','status']
    tabular_save(f'{prefix}_matches_and_errors',detail,columns=detail_cols)
    tabular_save(f'{prefix}_per_image',image_stats)
    xlsx_save(f'{prefix}.xlsx',{
        'Counts':(cols,rows_counts),'Row normalized':(cols,rows_norm),
        'Class metrics':(list(per_class[0]),per_class),
        'Matches and errors':(detail_cols,detail),
        'Per image':(list(image_stats[0]) if image_stats else ['image'],image_stats),
        'Summary':(['metric','value'],[{'metric':k,'value':v} for k,v in summary.items()])})
    json_save(f'{prefix}_summary.json',summary)
    plt = configure_plot()
    if nc>50:
        fig_w,fig_h,font_sz,left,bottom = 16,15,4.0,0.19,0.20
    else:
        fig_w,fig_h,font_sz,left,bottom = 13,12,7.0,0.12,0.15
    for suffix,arr in [('counts',matrix),('normalized',norm)]:
        fig,ax=plt.subplots(figsize=(fig_w,fig_h))
        ax.imshow(arr, cmap='Blues', interpolation='nearest', aspect='equal',vmin=0)
        ax.set_xticks(np.arange(nc), labels=names, rotation=90, fontsize=font_sz)
        ax.set_yticks(np.arange(nc), labels=names, fontsize=font_sz)
        ax.set_xlabel('Predicted label')
        ax.set_ylabel('True label')
        if nc==38:
            vmax=max(float(np.max(arr)),1e-10)
            for y in range(nc):
                for x in range(nc):
                    v=arr[y,x]
                    if v>0:
                        ax.text(x,y,str(int(v)) if suffix=='counts' else f'{v:.2f}',
                                ha='center',va='center',fontsize=4.0,
                                color='white' if v>.55*vmax else 'black')
        fig.subplots_adjust(left=left,bottom=bottom,right=0.965,top=0.96)
        save_figure(fig,Path(out_dir)/'figures'/f'{split}_{scope}_{suffix}')
        plt.close(fig)
    return summary


def write_training_epoch_tables(run_dir):
    p=Path(run_dir)/'ultralytics'/'results.csv'
    if not p.exists():
        raise FileNotFoundError(f'Ultralytics results.csv not found: {p}')
    with p.open('r',encoding='utf-8-sig',newline='') as f:
        rr=csv.DictReader(f)
        rows=[]
        for row in rr:
            clean={}
            for k,v in row.items():
                k=k.strip(); val=(v or '').strip()
                try:
                    num=float(val)
                    clean[k] = int(num) if k.lower()=='epoch' and num.is_integer() else num
                except ValueError:
                    clean[k]=val
            rows.append(clean)
    if not rows:
        raise ValueError('results.csv contains no epoch records')
    tabular_save(Path(run_dir)/'paper_outputs'/'training_epochs_numeric',rows)
    plt=configure_plot()
    x=[r.get('epoch',i+1) for i,r in enumerate(rows)]
    sets=[('Training_losses',['train/box_loss','train/cls_loss','train/dfl_loss']),
          ('Validation_losses',['val/box_loss','val/cls_loss','val/dfl_loss']),
          ('Detection_metrics',['metrics/precision(B)','metrics/recall(B)','metrics/mAP50(B)','metrics/mAP50-95(B)']),
          ('Learning_rates',['lr/pg0','lr/pg1','lr/pg2'])]
    for name, keys in sets:
        plotted=False
        fig,ax=plt.subplots(figsize=(7.2,5.2))
        for key in keys:
            if key in rows[0]:
                ax.plot(x,[r.get(key,np.nan) for r in rows],label=key.split('/')[-1],linewidth=1.8)
                plotted=True
        if plotted:
            ax.set_xlabel('Epoch')
            ax.set_ylabel('Value')
            ax.legend(frameon=False,fontsize=9)
            fig.subplots_adjust(left=0.14,right=0.96,top=0.95,bottom=0.14)
            save_figure(fig,Path(run_dir)/'paper_outputs'/'figures'/name)
        plt.close(fig)
    return rows


def export_native_metrics(native, names, split, out_dir):
                                              
    box=getattr(native,'box',None)
    ids=list(map(int,getattr(box,'ap_class_index',[]))) if box is not None else []
    index={cid:j for j,cid in enumerate(ids)}
    rows=[]
    for cid,name in enumerate(names):
        rec={'yolo_id':cid,'paper_number':cid+1 if cid<38 else None,
             'label_name':name,'in_validation':cid in index,
             'precision':None,'recall':None,'AP50':None,'AP50_95':None,'F1':None}
        if cid in index:
            vals=box.class_result(index[cid])
            p,r,ap50,ap5095=(float(x) for x in vals)
            rec.update(precision=p,recall=r,AP50=ap50,AP50_95=ap5095,
                       F1=2*p*r/(p+r) if p+r else 0.0)
        rows.append(rec)
    tabular_save(Path(out_dir)/f'{split}_native_105_per_class',rows)
    tabular_save(Path(out_dir)/f'{split}_native_numeric38_per_class',rows[:38])
    speed={k:float(v) for k,v in (getattr(native,'speed',None) or {}).items() if isinstance(v,(int,float))}
                                       
    meta={'split':split,'official_metrics':{k:float(v) for k,v in getattr(native,'results_dict',{}).items()
                                             if isinstance(v,(int,float))},'ms_per_image':speed,
          'fps_inference_only':1000/speed['inference'] if speed.get('inference',0)>0 else None,
          'FPS_note':'Calculated as 1000/inference_ms only; excludes preprocessing, postprocessing, and data loading, so it is not end-to-end FPS'}
    json_save(Path(out_dir)/f'{split}_official_metrics.json',meta)
    tabular_save(Path(out_dir)/f'{split}_official_metrics',[
        {'metric':k,'value':v} for k,v in meta['official_metrics'].items()]+[
        {'metric':'speed_ms_'+k,'value':v} for k,v in speed.items()]+[
        {'metric':'fps_inference_only','value':meta['fps_inference_only']}])
    return meta


def export_native_curves(native, names, split, out_dir):
    ''
    box=getattr(native,'box',None)
    candidate=[]
    try:
        result=getattr(native,'curves_results',None)
        if result is None and box is not None:
            result=getattr(box,'curves_results',None)
        candidate=list(result) if result is not None else []
    except Exception:
        candidate=[]
    names_curve=getattr(native,'curves',None)
    if not names_curve and box is not None:
        names_curve=getattr(box,'curves',None)
    names_curve=list(names_curve or [])
    warnings=[]
    saved=[]
    for i, rec in enumerate(candidate):
        try:
            if len(rec)<2:
                continue
            xx=np.asarray(rec[0],dtype=float).reshape(-1)
            yy=np.asarray(rec[1],dtype=float)
            if yy.ndim==1:
                yy=yy[None,:]
            if yy.shape[-1] != len(xx) and yy.shape[0]==len(xx):
                yy=yy.T
            if yy.shape[-1]!=len(xx):
                warnings.append(f'Curve {i}: dimensions {yy.shape}do not match X-axis size{len(xx)}')
                continue
            curve_name=str(names_curve[i] if i<len(names_curve) else f'curve_{i+1}')
                                               
            class_ids=list(map(int,getattr(box,'ap_class_index',[]))) if box is not None else []
            class_labels=([names[j] for j in class_ids] if len(class_ids)==yy.shape[0]
                          else names if yy.shape[0]==len(names)
                          else [f'series_{j}' for j in range(yy.shape[0])])
            keys=['x']+[f'c{j:03d}_{name}' for j,name in enumerate(class_labels)]
            rows=[{'x':float(xx[k]),**{keys[j+1]:float(yy[j,k]) for j in range(yy.shape[0])}}
                  for k in range(len(xx))]
            stem=f'{split}_curve_{i+1}_{"".join(c if c.isalnum() else "_" for c in curve_name)[:35]}'
            tabular_save(Path(out_dir)/'curve_tables'/stem,rows,keys)
            json_save(Path(out_dir)/'curve_tables'/f'{stem}_metadata.json',{
                'curve_name':curve_name,'x_label':str(rec[2]) if len(rec)>2 else '',
                'y_label':str(rec[3]) if len(rec)>3 else '',
                'series_count':int(yy.shape[0]),'series_names':class_labels,
                'source':'native Ultralytics DetMetrics.curves_results'})
            plt=configure_plot()
            fig,ax=plt.subplots(figsize=(7.4,5.4))
            if yy.shape[0]>15:
                avg=np.mean(yy,axis=0)
                ax.plot(xx,avg,linewidth=2,label='Mean of available classes')
            else:
                for j in range(yy.shape[0]):
                    ax.plot(xx,yy[j],linewidth=1,label=str(class_labels[j]))
                ax.legend(frameon=False,fontsize=7,ncol=2)
            ax.set_xlabel(str(rec[2]) if len(rec)>2 else 'X')
            ax.set_ylabel(str(rec[3]) if len(rec)>3 else 'Y')
            fig.subplots_adjust(left=.14,right=.96,bottom=.14,top=.95)
            save_figure(fig,Path(out_dir)/'figures'/stem)
            plt.close(fig)
            saved.append(stem)
        except Exception as exc:
            warnings.append(f'Curve {i}: {exc}')
    if not saved:
        warnings.append('The installed Ultralytics version did not provide exportable PR/P/R/F1 curve coordinates. Native plots may be present; no curve values were fabricated.')
    json_save(Path(out_dir)/f'{split}_curve_export_status.json',{'saved':saved,'warnings':warnings})
    return {'saved':saved,'warnings':warnings}


def infer_and_export(model, split, imgdir, names, output_dir, imgsz, device,
                     conf_export=0.01, cm_conf=0.25, nms_iou=0.7, match_iou=0.50, save_visuals=False):
                                                     
    if conf_export > cm_conf:
        raise ValueError('Prediction export confidence threshold must be <= confusion matrix confidence threshold')
    images=sorted(p for p in imgdir.rglob('*') if p.is_file() and p.suffix.lower() in IMAGE_EXT)
    all_gt, all_predictions, gt_rows, pred_rows, image_rows = {},{},{},{},[]
    export_gt=[];export_pred=[]
    for idx,img in enumerate(images,1):
        gt,w,h=gt_boxes(img)
        result=model.predict(source=str(img),imgsz=imgsz,device=device,conf=conf_export,
                             iou=nms_iou,verbose=False,save=False,max_det=300)[0]
        if save_visuals:
            from PIL import Image as PILImage
            visual_dir=Path(output_dir)/'predicted_images'
            visual_dir.mkdir(parents=True,exist_ok=True)
                                                  
            drawn=result.plot()
            PILImage.fromarray(drawn[...,::-1]).save(visual_dir/(img.stem+'.png'))
        pred=[]
        if result.boxes is not None:
            for j,b in enumerate(result.boxes):
                cls=int(b.cls.item())
                if not 0<=cls<105:
                    raise ValueError(f'Invalid predicted class: {cls}')
                pred.append({'id':cls,'box':[float(x) for x in b.xyxy[0].tolist()],
                             'confidence':float(b.conf.item()),'index':j})
        key=str(img)
        all_gt[key]=gt
        all_predictions[key]=pred
        image_rows.append({'split':split,'image':key,'width':w,'height':h,'gt_all':len(gt),
                           'gt_numeric':sum(g['id']<38 for g in gt),
                           'pred_all_exported':len(pred),
                           'pred_numeric_exported':sum(p['id']<38 for p in pred),
                           'cm_conf':cm_conf,'export_conf':conf_export})
        for g in gt:
            export_gt.append({'split':split,'image':key,'index':g['index'],'yolo_id':g['id'],
                              'paper_number':g['id']+1 if g['id']<38 else None,
                              'label':names[g['id']],
                              **dict(zip(('x1','y1','x2','y2'),g['box']))})
        for p in pred:
            export_pred.append({'split':split,'image':key,'index':p['index'],'yolo_id':p['id'],
                                'paper_number':p['id']+1 if p['id']<38 else None,
                                'label':names[p['id']],'confidence':p['confidence'],
                                **dict(zip(('x1','y1','x2','y2'),p['box']))})
        if idx%40==0 or idx==len(images):
            print(f'[EVAL] {split} {idx}/{len(images)}',flush=True)
    folder=Path(output_dir)
    tabular_save(folder/f'{split}_all_GT_boxes',export_gt)
    tabular_save(folder/f'{split}_all_predictions_conf_ge_{str(conf_export).replace(".","p")}',export_pred,
                 list(export_pred[0]) if export_pred else ['split','image','index','yolo_id','confidence'])
    tabular_save(folder/f'{split}_per_image',image_rows)
    summaries=[]
    for scope,scope_names in [('full105',names),('numeric38',names[:38])]:
        summaries.append(matrix_report(all_gt,all_predictions,scope_names,split,scope,folder/'double_confusion',cm_conf,match_iou))
    tabular_save(folder/'double_confusion'/f'{split}_matrix_comparison',summaries)
    return summaries


def initial_stats(net, run_dir, input_imgsz, stage):
    stat={'stage':stage,'module_count':len(list(net.modules())),
          'parameter_count':sum(p.numel() for p in net.parameters()),
          'trainable_count':sum(p.numel() for p in net.parameters() if p.requires_grad),
          'imgsz_training':input_imgsz,
          'strides':list(map(float,net.stride)),
          'class_count_at_this_stage':int(getattr(net.model[-1],'nc',-1))}
    try:
        from ultralytics.utils.torch_utils import get_flops
        stat['GFLOPs_640']=float(get_flops(net,imgsz=640))
        stat[f'GFLOPs_{input_imgsz}']=float(get_flops(net,imgsz=input_imgsz))
    except Exception as e:
        stat['flops_error']=str(e)
    tabular_save(Path(run_dir)/f'complexity_{stage}',[{'metric':k,'value':v} for k,v in stat.items()])
    return stat


def build_model(opt, pretrained):
    from ultralytics import YOLO
    if opt:
        from reconstructed_yolo11m import register_model_modules,enable_wiou_v3
        register_model_modules()
        if USE_WIOU:
            enable_wiou_v3()
        mpath=ROOT/'models'/'yolo11m_opt.yaml'
        if not mpath.is_file():
            raise FileNotFoundError(f'Missing custom model YAML: {mpath}')
        model=YOLO(str(mpath),task='detect')
        if [float(a) for a in model.model.stride]!=[4.0,8.0,16.0,32.0]:
            raise RuntimeError('Modified model P2/P3/P4/P5 strides must be 4/8/16/32')
        names=Counter(type(m).__name__ for m in model.model.modules())
        for k in ('SPDConv','DySample','PaperCBAM'):
            if names[k]<1:
                raise RuntimeError(f'Modified network missing {k}')
        if pretrained:
            w=ROOT/'yolo11m.pt'
            if not w.is_file():
                raise FileNotFoundError('Pretrained weights were requested but yolo11m.pt was not found in the project root')
            model.load(str(w))
    else:
        if pretrained:
            w=ROOT/'yolo11m.pt'
            if not w.exists():
                raise FileNotFoundError('Pretrained weights were requested but yolo11m.pt was not found in the project root')
            model=YOLO(str(w),task='detect')
        else:
            model=YOLO('yolo11m.yaml',task='detect')
        names=Counter(type(m).__name__ for m in model.model.modules())
        if any(names[k] for k in ('SPDConv','DySample','PaperCBAM')):
            raise RuntimeError('Baseline unexpectedly contains custom modules')
    return model


def make_manifest(run_dir):
    run_dir=Path(run_dir)
    rows=[]
    for p in sorted(run_dir.rglob('*')):
        if p.is_file() and '_file_audit' not in p.parts and p.name!='output_manifest.xlsx':
            rel=str(p.relative_to(run_dir))
            rows.append({'file':rel,'bytes':p.stat().st_size,
                         'sha256':sha256(p) if p.suffix.lower() in {'.csv','.xlsx','.json','.yaml','.png','.tiff','.txt'} else None})
    tabular_save(run_dir/'paper_outputs'/'output_manifest',rows)
    return rows

def verify_final_outputs(run_dir, splits):
    ''
    from PIL import Image
    from openpyxl import load_workbook
    root = Path(run_dir)
    out = root/'paper_outputs'
    checks=[]
    must=[root/'ultralytics'/'weights'/'best.pt',
          root/'ultralytics'/'weights'/'last.pt',
          root/'ultralytics'/'results.csv',
          out/'training_epochs_numeric.csv',
          out/'training_epochs_numeric.xlsx',
          root/'complexity_before_training.xlsx',
          root/'complexity_trained_fused.xlsx',
          out/'dataset_audit'/'class_mapping_105.xlsx',
          out/'dataset_audit'/'dataset_per_image.xlsx']
    for split in splits:
        must.extend([out/'native_metrics'/f'{split}_official_metrics.xlsx',
                     out/'native_metrics'/f'{split}_native_105_per_class.xlsx',
                     out/'native_metrics'/f'{split}_native_numeric38_per_class.xlsx',
                     out/f'{split}_all_GT_boxes.xlsx',
                     out/f'{split}_per_image.xlsx',
                     out/'double_confusion'/f'{split}_matrix_comparison.xlsx'])
        for scope,n in [('full105',105),('numeric38',38)]:
            prefix = out/'double_confusion'/f'{split}_{scope}'
            must.extend([prefix.with_suffix('.xlsx'),
                         Path(str(prefix)+'_counts.csv'),
                         Path(str(prefix)+'_normalized.csv'),
                         Path(str(prefix)+'_per_class.xlsx'),
                         Path(str(prefix)+'_matches_and_errors.xlsx')])
            for suffix in ('counts','normalized'):
                figbase=out/'double_confusion'/'figures'/f'{split}_{scope}_{suffix}'
                must.extend([Path(str(figbase)+'.png'),Path(str(figbase)+'.tiff')])
    for p in must:
        status='OK' if p.is_file() and p.stat().st_size>0 else 'MISSING'
        note=''
        if status=='OK' and p.suffix.lower() in ('.png','.tiff'):
            try:
                with Image.open(p) as im:
                    dpi=im.info.get('dpi')
                    if dpi is None or min(map(float,dpi)) < (299 if p.suffix.lower()=='.png' else 599):
                        status='INVALID_DPI'
                    note=f'dpi={dpi}; px={im.size}'
            except Exception as exc:
                status='BAD_IMAGE';note=str(exc)
        elif status=='OK' and p.suffix.lower()=='.xlsx' and '_full105.xlsx' in p.name or (status=='OK' and p.name.endswith('_numeric38.xlsx') and p.parent.name=='double_confusion'):
            try:
                nc=105 if 'full105' in p.name else 38
                wb=load_workbook(p,read_only=True,data_only=True)
                ws=wb['Counts']
                if ws.max_row!=nc+1 or ws.max_column!=nc+1:
                    status='BAD_MATRIX_SHAPE'
                note=f'expected={(nc+1,nc+1)}, actual={(ws.max_row,ws.max_column)}'
                wb.close()
            except Exception as exc:
                status='BAD_MATRIX';note=str(exc)
        checks.append({'path':str(p.relative_to(root)),'status':status,
                       'bytes':p.stat().st_size if p.exists() else 0,'note':note})
    tabular_save(out/'export_audit',checks)
    errors=[x for x in checks if x['status']!='OK']
    json_save(out/'export_audit_summary.json',{'required_checks':len(checks),
                 'errors':errors,'success':not errors})
    print(f'[EXPORT AUDIT] Checked {len(checks)} required files, issues: {len(errors)}')
    if errors:
        raise RuntimeError('Paper outputs are incomplete; see paper_outputs/export_audit.xlsx: '+str(errors[:5]))
    return checks


def run_single_split_original():
    import torch
    from ultralytics import YOLO
    os.chdir(ROOT)
    if DEVICE!='cpu' and not torch.cuda.is_available():
        raise RuntimeError('CUDA is unavailable; check the song environment first')
    data_yaml,names,split_dirs=read_data(DATA_YAML)
    stamp=dt.datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    run_dir=ROOT/'runs'/'paper_complete_experiments'/f'{stamp}_{MODEL_TAG}'
    run_dir.mkdir(parents=True,exist_ok=False)
    json_save(run_dir/'status.json',{'status':'running','time':stamp})
    try:
        audits=audit_dataset(names,split_dirs,run_dir/'paper_outputs'/'dataset_audit')
        shutil.copy2(data_yaml,run_dir/'dataset_snapshot.yaml')
        if OPT_MODEL:
            mpath=ROOT/'models'/'yolo11m_opt.yaml'
            if mpath.is_file():
                shutil.copy2(mpath,run_dir/'model_snapshot.yaml')
        config={'model':MODEL_TAG,'model_kind':'YOLO11m-opt reconstructed' if OPT_MODEL else 'official YOLO11m baseline',
                'data_yaml':str(data_yaml),'data_sha256':sha256(data_yaml),'num_classes':105,
                'numeric_yolo_ids_0based':list(range(38)),'paper_numeric_labels_1based':list(range(1,39)),
                'epochs':EPOCHS,'imgsz':IMGSZ,'batch':BATCH,'optimizer':OPTIMIZER,
                'device':DEVICE,'use_pretrained':USE_PRETRAINED,'use_wiou':bool(OPT_MODEL and USE_WIOU),
                'evaluate_test':EVALUATE_TEST_AFTER_TRAINING,
                'search_origin':'Trial8 best among 10 completed tuning trials' if OPT_MODEL else 'Ultralytics defaults (no custom tuning)',
                'preflight':audits,'confusion_conf':CM_CONF,'prediction_export_conf':EXPORT_CONF,
                'match_iou':MATCH_IOU,'nms_iou':NMS_IOU,'max_det_predictions':300,'seed':SEED,
                'fairness_note':'Baseline640 vs Opt960; cannot attribute entire performance difference to architecture alone.'}
        if OPT_MODEL:
            from train_rebuilt import PAPER_HYP
            final_hyp={**PAPER_HYP,**TRIAL8_BEST_HYP}
            config['full_hyperparameters']=final_hyp
        json_save(run_dir/'run_config.json',config)
        import platform
        json_save(run_dir/'environment.json',{'python':sys.version,'platform':platform.platform(),
                  'torch':torch.__version__,'ultralytics':__import__('ultralytics').__version__,
                  'cuda':torch.version.cuda,'gpu':torch.cuda.get_device_name(0) if torch.cuda.is_available() else None})
        model=build_model(OPT_MODEL,USE_PRETRAINED)
        before=initial_stats(model.model,run_dir,IMGSZ,'before_training')
        args={'data':str(data_yaml),'epochs':EPOCHS,'imgsz':IMGSZ,'batch':BATCH,'device':DEVICE,
              'workers':WORKERS,'seed':SEED,'deterministic':False,'amp':AMP,
              'cache':False,'save':True,'save_period':SAVE_PERIOD,'val':True,
              'plots':True,'verbose':True,'patience':0,'close_mosaic':50 if OPT_MODEL else 10,
              'project':str(run_dir),'name':'ultralytics','exist_ok':False,
              'pretrained':USE_PRETRAINED}
        if OPT_MODEL:
            args.update({'optimizer':'SGD','nbs':8})
            args.update(final_hyp)
        else:
            args['optimizer']='auto'
        json_save(run_dir/'train_arguments.json',args)
                                      
        epoch_jsonl = run_dir/'paper_outputs'/'epoch_live.jsonl'
        epoch_jsonl.parent.mkdir(parents=True,exist_ok=True)
        def _record_epoch(trainer):
            rec={'epoch':int(trainer.epoch)+1,
                 'timestamp':dt.datetime.now().astimezone().isoformat()}
            for k,v in (getattr(trainer,'metrics',None) or {}).items():
                try:
                    v=float(v)
                    if math.isfinite(v):
                        rec[k]=v
                except (TypeError,ValueError):
                    pass
            try:
                rec['learning_rates']=[float(g['lr']) for g in trainer.optimizer.param_groups]
            except Exception:
                pass
            with epoch_jsonl.open('a',encoding='utf-8') as f:
                f.write(json.dumps(rec,ensure_ascii=False)+'\n')
                f.flush()
        model.add_callback('on_fit_epoch_end',_record_epoch)
                                    
        import logging
        from ultralytics.utils import LOGGER
        file_handler=logging.FileHandler(run_dir/'ultralytics_training.log',encoding='utf-8')
        file_handler.setLevel(logging.INFO)
        LOGGER.addHandler(file_handler)
        print('[TRAIN]',MODEL_TAG,'epochs',EPOCHS,'imgsz',IMGSZ,'batch',BATCH,'use_wiou',OPT_MODEL and USE_WIOU)
        model.train(**args)
        result_dir=run_dir/'ultralytics'
        best=result_dir/'weights'/'best.pt'
        last=result_dir/'weights'/'last.pt'
        if not best.is_file() or not last.is_file():
            raise RuntimeError('best.pt or last.pt was not generated; see the original training log')
        rows=write_training_epoch_tables(run_dir)
        if OPT_MODEL:
            from reconstructed_yolo11m import register_model_modules
            register_model_modules()
        trained=YOLO(str(best),task='detect')
        det=trained.model.model[-1]
        if int(getattr(det,'nc',-1))!=105:
            raise RuntimeError(f'Trained model nc={getattr(det,"nc",None)}, expected 105')
        trained_names=trained.model.names
        if isinstance(trained_names,dict):
            trained_names=[trained_names.get(i,trained_names.get(str(i))) for i in range(105)]
        if [str(x) for x in trained_names]!=names:
            raise RuntimeError('best.pt class-name ordering does not match data.yaml')
        complexity=initial_stats(trained.model,run_dir,IMGSZ,'trained_fused')
        json_save(run_dir/'model_complexity_comparison.json',{'initial':before,'after_best_fused':complexity,
                  'best_weight_size_MB':best.stat().st_size/1048576,
                  'note':'Module and parameter counts before and after fusion are not directly comparable; use the same model state in manuscript comparisons'})
        splits=['val']+(['test'] if EVALUATE_TEST_AFTER_TRAINING else [])
        summary_all=[]
        for split in splits:
            print('[EVAL] Native Ultralytics metrics',split)
            native=trained.val(data=str(data_yaml),split=split,imgsz=IMGSZ,batch=BATCH,device=DEVICE,
                               workers=WORKERS,plots=True,verbose=False)
            native_out=run_dir/'paper_outputs'/'native_metrics'
            metrics=export_native_metrics(native,names,split,native_out)
            curves=export_native_curves(native,names,split,run_dir/'paper_outputs')
            pred_summaries=infer_and_export(trained,split,split_dirs[split],names,
                                            run_dir/'paper_outputs',IMGSZ,DEVICE,
                                            conf_export=EXPORT_CONF,cm_conf=CM_CONF,
                                            nms_iou=NMS_IOU,match_iou=MATCH_IOU)
            for s in pred_summaries:
                summary_all.append({'split':split,**s,'official_mAP50_95':metrics['official_metrics'].get('metrics/mAP50-95(B)')})
            if curves['warnings']:
                print('[WARNING] Native curve tables',curves['warnings'])
        tabular_save(run_dir/'paper_outputs'/'metrics_and_confusion_summary',summary_all)
        verification=verify_final_outputs(run_dir,splits)
        filelist=make_manifest(run_dir)
        json_save(run_dir/'status.json',{'status':'completed','best_pt':str(best),'last_pt':str(last),
                  'eval_splits':splits,'files_count':len(filelist),
                  'completed_time':dt.datetime.now().isoformat()})
        print('\n[FINISHED]',run_dir)
        print('best.pt',best)
        print('105x105 and 38x38 confusion matrices and source CSV/XLSX:',run_dir/'paper_outputs'/'double_confusion')
        print('Original PR/F1/P/R curves:',run_dir/'paper_outputs'/'curve_tables')
        print('Complete per-epoch table:',run_dir/'paper_outputs'/'training_epochs_numeric.xlsx')
        return best
    except BaseException as exc:
        json_save(run_dir/'status.json',{'status':'interrupted' if isinstance(exc,KeyboardInterrupt) else 'failed',
                  'error':repr(exc),'time':dt.datetime.now().isoformat()})
        (run_dir/'error.log').write_text(traceback.format_exc(),encoding='utf-8')
        raise
    finally:
        if 'file_handler' in locals():
            try:
                from ultralytics.utils import LOGGER
                LOGGER.removeHandler(file_handler)
                file_handler.close()
            except Exception:
                pass

def _external_prepare(images_dir, labels_dir, outdir, names):
    ''
    import yaml
    imgroot=Path(images_dir).expanduser().resolve()
    labroot=Path(labels_dir).expanduser().resolve()
    if not imgroot.is_dir():
        raise FileNotFoundError(f'EXTERNAL_IMAGES directory not found: {imgroot}')
    if not labroot.is_dir():
        raise FileNotFoundError(f'EXTERNAL_LABELS directory not found: {labroot}')
    images=sorted(p for p in imgroot.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXT)
    if not images:
        raise RuntimeError(f'External images directory is empty: {imgroot}')
    names_lower=set()
    for p in images:
        if p.name.casefold() in names_lower:
            raise RuntimeError(f'Case-insensitive image filename collision: {p.name}')
        names_lower.add(p.name.casefold())
    staged=Path(outdir)/'external_dataset_snapshot'
    imgout=staged/'images'; labout=staged/'labels'
    imgout.mkdir(parents=True,exist_ok=True)
    labout.mkdir(parents=True,exist_ok=True)
    rows=[]
    for p in images:
        lp=labroot/(p.stem+'.txt')
        if not lp.is_file() and not ALLOW_MISSING_EXTERNAL_LABELS_AS_EMPTY:
            raise FileNotFoundError(f'External image is missing matching TXT annotation: {lp}')
        target_img=imgout/p.name
        target_lab=labout/(p.stem+'.txt')
        if target_img.exists() or target_lab.exists():
            raise FileExistsError(f'Destination already exists in staging area; change output directory: {target_img}')
        try:
            os.link(p,target_img)
        except OSError:
            shutil.copy2(p,target_img)
        if lp.is_file():
            shutil.copy2(lp,target_lab)
        else:
            target_lab.write_text('',encoding='utf-8')
                                         
        gt, w, h=gt_boxes(target_img)
        rows.append({'image':p.name,'original_image_path':str(p),
                     'original_label_path':str(lp) if lp.is_file() else '',
                     'width':w,'height':h,'gt_all':len(gt),
                     'gt_numeric38':sum(g['id']<38 for g in gt),
                     'has_label_file':lp.is_file()})
    tabular_save(Path(outdir)/'paper_outputs'/'external_source_manifest',rows)
    yaml_path=Path(outdir)/'external_eval_data.yaml'
    cfg={'path':str(staged.resolve()),'train':'images','val':'images','test':'images',
         'nc':105,'names':names}
    yaml_path.write_text(yaml.safe_dump(cfg,sort_keys=False,allow_unicode=True),encoding='utf-8')
    return yaml_path,imgout,rows


def _check_external_outputs(run_dir):
    from openpyxl import load_workbook
    root=Path(run_dir)/'paper_outputs'
    must=[root/'external_source_manifest.xlsx',
          root/'native_metrics'/'external_native_105_per_class.xlsx',
          root/'native_metrics'/'external_native_numeric38_per_class.xlsx',
          root/'external_all_GT_boxes.xlsx',root/'external_per_image.xlsx',
          root/'double_confusion'/'external_full105.xlsx',
          root/'double_confusion'/'external_numeric38.xlsx']
    for scope,n in [('full105',105),('numeric38',38)]:
        stem=root/'double_confusion'/f'external_{scope}'
        must += [Path(str(stem)+'_counts.csv'),Path(str(stem)+'_normalized.csv'),
                 Path(str(stem)+'_matches_and_errors.xlsx'),Path(str(stem)+'_per_class.xlsx')]
        for suffix in ('counts','normalized'):
            fig=root/'double_confusion'/'figures'/f'external_{scope}_{suffix}'
            must += [Path(str(fig)+'.png'),Path(str(fig)+'.tiff')]
    checks=[]
    for p in must:
        status='OK' if p.is_file() and p.stat().st_size>0 else 'MISSING'
        if status=='OK' and p.name in ('external_full105.xlsx','external_numeric38.xlsx'):
            try:
                n=105 if 'full105' in p.name else 38
                wb=load_workbook(p,read_only=True,data_only=True)
                ws=wb['Counts']
                if ws.max_row!=n+1 or ws.max_column!=n+1: status='BAD_MATRIX_SHAPE'
                wb.close()
            except Exception as e:
                status='BAD_MATRIX_'+str(e)
        checks.append({'file':str(p.relative_to(run_dir)),'status':status})
    tabular_save(root/'external_export_audit',checks)
    errors=[r for r in checks if r['status']!='OK']
    if errors:
        raise RuntimeError('External evaluation is missing required outputs: '+repr(errors[:8]))
    print('[AUDIT] Required external-evaluation files validated:',len(checks))


def run_external_only(weight_path=None):
    import torch
    from ultralytics import YOLO
    os.chdir(ROOT)
    if DEVICE!='cpu' and not torch.cuda.is_available():
        raise RuntimeError('CUDA is unavailable; check the song environment')
    model_path=Path(weight_path or EXISTING_BEST_PT).expanduser().resolve()
    if not model_path.is_file():
        raise FileNotFoundError(f'best.pt not found; update EXISTING_BEST_PT: {model_path}')
    data_yaml,names,_=read_data(DATA_YAML)
    if OPT_MODEL:
        from reconstructed_yolo11m import register_model_modules
        register_model_modules()
                                        
    model=YOLO(str(model_path),task='detect')
    if int(getattr(model.model.model[-1],'nc',-1))!=105:
        raise ValueError('Checkpoint is not a 105-class detector; check the best.pt path')
    loaded_names=model.names
    got=[str(loaded_names.get(i,loaded_names.get(str(i))) if isinstance(loaded_names,dict)
             else loaded_names[i]) for i in range(105)]
    if got!=names:
        raise ValueError('The 105-class order in best.pt differs from DATA_YAML; verify the mapping before comparison')
    stamp=dt.datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    run_dir=ROOT/'runs'/'external_evaluation'/f'{stamp}_{MODEL_TAG}'
    run_dir.mkdir(parents=True,exist_ok=False)
    json_save(run_dir/'status.json',{'status':'running','mode':RUN_MODE})
    try:
        shutil.copy2(data_yaml,run_dir/'original_data_snapshot.yaml')
        cfg, images, manifest=_external_prepare(EXTERNAL_IMAGES,EXTERNAL_LABELS,run_dir,names)
        config={'mode':RUN_MODE,'weights':str(model_path),'weights_sha256':sha256(model_path),
                'source_data_yaml':str(data_yaml),'external_images':EXTERNAL_IMAGES,
                'external_labels':EXTERNAL_LABELS,'images_count':len(manifest),
                'imgsz':IMGSZ,'batch':BATCH,'device':DEVICE,
                'export_conf':EXPORT_CONF,'matrix_conf':CM_CONF,'match_iou':MATCH_IOU,
                'nms_iou':NMS_IOU,'scope':'105 labels + terminal numeric 38 labels',
                'note':'Ground-truth labels are required to calculate precision, recall, and AP; images alone are insufficient for quantitative evaluation.'}
        json_save(run_dir/'run_config.json',config)
        initial_stats(model.model,run_dir,IMGSZ,'loaded_for_external_eval')
        print(f'[EXTERNAL] Loaded weights: {model_path}')
        print(f'[EXTERNAL] Number of images: {len(manifest)}; evaluating using the selected best.pt (no retraining)')
                                                 
        native=model.val(data=str(cfg),split='val',imgsz=IMGSZ,batch=BATCH,device=DEVICE,
                         workers=WORKERS,plots=True,verbose=False,conf=EXPORT_CONF,
                         iou=NMS_IOU,project=str(run_dir),name='ultralytics_external',exist_ok=True)
        out=run_dir/'paper_outputs'
        export_native_metrics(native,names,'external',out/'native_metrics')
        export_native_curves(native,names,'external',out)
        infer_and_export(model,'external',images,names,out,IMGSZ,DEVICE,
                         conf_export=EXPORT_CONF,cm_conf=CM_CONF,
                         nms_iou=NMS_IOU,match_iou=MATCH_IOU,
                         save_visuals=SAVE_EXTERNAL_VISUALS)
        _check_external_outputs(run_dir)
        make_manifest(run_dir)
        json_save(run_dir/'status.json',{'status':'completed','mode':RUN_MODE,
                                        'weights':str(model_path),'image_count':len(manifest)})
        print('\n[EXTERNAL DONE] Results saved in:',run_dir)
        print('105-class metrics:',out/'native_metrics'/'external_native_105_per_class.xlsx')
        print('38-class metrics:',out/'native_metrics'/'external_native_numeric38_per_class.xlsx')
        print('Confusion matrices and figures:',out/'double_confusion')
        print('Prediction images:',out/'predicted_images')
        return run_dir
    except BaseException as exc:
        json_save(run_dir/'status.json',{'status':'failed','error':repr(exc)})
        (run_dir/'error.log').write_text(traceback.format_exc(),encoding='utf-8')
        raise


def run():
    if RUN_MODE=='external_only':
        return run_external_only()
    if RUN_MODE=='train':
        return run_single_split_original()
    if RUN_MODE=='train_then_external':
                                                
        ext_images=Path(EXTERNAL_IMAGES)
        ext_labels=Path(EXTERNAL_LABELS)
        if not ext_images.is_dir() or not ext_labels.is_dir():
            raise FileNotFoundError('External images/labels directory does not exist; verify EXTERNAL_IMAGES and EXTERNAL_LABELS')
        external_images=[p for p in ext_images.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXT]
        if not external_images:
            raise ValueError(f'External images directory contains no images: {ext_images}')
        if not ALLOW_MISSING_EXTERNAL_LABELS_AS_EMPTY:
            missing=[p.name for p in external_images if not (ext_labels/(p.stem+'.txt')).is_file()]
            if missing:
                raise FileNotFoundError(f'External images without matching YOLO TXT labels: {len(missing)} images, examples: {missing[:5]}')
        print('[PIPELINE] External dataset checked. Training will start, followed by evaluation with the newly trained best.pt.',flush=True)
        best=run_single_split_original()
        if not best or not Path(best).is_file():
            raise RuntimeError('Training did not produce best.pt; external evaluation cannot start')
        print(f'[PIPELINE] Newly trained best.pt: {best}',flush=True)
        return run_external_only(best)
    raise ValueError('RUN_MODE must be external_only, train, or train_then_external')


if __name__=='__main__':
    run()
